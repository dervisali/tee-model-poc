"""
Çoklu format belge loader'ı — Phase 1.

PDF, DOCX, PPTX, JPEG/PNG dosyalarından sayfa-bazlı metin çıkarır. Her loader
`list[PageContent]` döndürür; `iter_documents` bir kök dizini özyinelemeli
gezer ve her dosya için sonuç yield eder.

Bağımlılıklar (pdfplumber, python-docx, python-pptx, pytesseract, Pillow)
loader fonksiyonu çağrıldığında lazy import edilir; biri eksikse modülün
geri kalanı yine import edilebilir.

PDF OCR fallback dahildir: pdfplumber boş döndürdüğünde sayfa pypdfium2
ile imaja render edilip pytesseract ile OCR'lanır. tesseract sistem ikilisi
yoksa fallback sessizce devre dışı kalır ve sayfa boş raporlanır.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Iterator, NamedTuple


logger = logging.getLogger(__name__)


SUPPORTED_EXTENSIONS: frozenset[str] = frozenset({".pdf", ".docx", ".pptx", ".jpeg", ".jpg", ".png"})
IMAGE_EXTENSIONS: frozenset[str] = frozenset({".jpeg", ".jpg", ".png"})


class PageContent(NamedTuple):
    """Tek bir sayfa / slide / görsel için çıkarılan metin."""

    text: str
    page: int | None         # PDF sayfa numarası, PPTX slide numarası; DOCX/imaj için None
    total_pages: int | None  # Belgenin toplam sayfa sayısı


class LoadResult(NamedTuple):
    """`iter_documents` için tek dosya kaydı."""

    path: Path
    pages: list[PageContent]   # Hata durumunda boş liste
    error: str | None          # Başarıda None


class LoaderError(Exception):
    """Bir loader bağımlılığı eksik veya extraction başarısız olduğunda."""


# ---------------------------------------------------------------------------
# Yardımcılar
# ---------------------------------------------------------------------------

def _normalize(text: str) -> str:
    """Tek-boşluk normalizasyonu; embedding ve BM25 için tutarlı girdi."""
    return " ".join(text.split())


def _tesseract_runtime_errors(pytesseract_module) -> tuple[type[BaseException], ...]:
    """pytesseract sürümlerinde değişebilen OCR runtime hatalarını yakala."""
    error_types: list[type[BaseException]] = [UnicodeDecodeError]
    for name in ("TesseractNotFoundError", "TesseractError"):
        exc_type = getattr(pytesseract_module, name, None)
        if isinstance(exc_type, type) and issubclass(exc_type, BaseException):
            error_types.append(exc_type)
    return tuple(error_types)


def _tesseract_image_to_string(image, *, lang: str, source_path: "Path | None" = None) -> str:
    """Bir PIL imajını OCR'lar; pytesseract'ın stderr decode hatasına dayanıklı.

    pytesseract'ın `get_errors()` fonksiyonu tesseract ikilisinin stderr'ini
    katı UTF-8 ile çözer ve bazı tesseract sürümlerinin stderr'e yazdığı
    UTF-8 olmayan baytlarda `UnicodeDecodeError` yükseltir — oysa OCR aslında
    başarılıdır (stdout metni geçerli UTF-8'dir). Bu durumda ikiliyi doğrudan
    çağırıp stdout'u okuruz. `source_path` verilirse (imaj dosyaları) ikili
    orijinal dosya üzerinde çalışır; geçici dosya gerekmez.
    """
    import pytesseract

    try:
        return pytesseract.image_to_string(image, lang=lang)
    except UnicodeDecodeError:
        return _tesseract_via_subprocess(image, lang=lang, source_path=source_path)


def _tesseract_via_subprocess(image, *, lang: str, source_path: "Path | None" = None) -> str:
    """tesseract ikilisini doğrudan çalıştırır; stdout/stderr'i hoşgörüyle çözer."""
    import os
    import shutil
    import subprocess
    import tempfile

    binary = shutil.which("tesseract") or "tesseract"
    tmp_path: str | None = None
    if source_path is not None:
        target = str(source_path)
    else:
        fd, tmp_path = tempfile.mkstemp(suffix=".png")
        os.close(fd)
        image.save(tmp_path, format="PNG")
        target = tmp_path
    try:
        proc = subprocess.run(
            [binary, target, "stdout", "-l", lang],
            capture_output=True,
            timeout=90,
        )
        # tesseract sıfır-olmayan çıkış kodunda bile çoğu zaman stdout'a
        # kullanılabilir metin yazar. stdout'u hoşgörüyle çöz; stderr'i asla
        # katı çözme.
        text = proc.stdout.decode("utf-8", errors="replace")
        if proc.returncode != 0 and not text.strip():
            stderr = proc.stderr.decode("utf-8", errors="replace").strip()
            detail = f": {stderr}" if stderr else ""
            raise LoaderError(f"OCR başarısız: tesseract exit {proc.returncode}{detail}")
        return text
    except subprocess.TimeoutExpired as exc:
        raise LoaderError("OCR başarısız: tesseract timeout") from exc
    finally:
        if tmp_path is not None:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


# ---------------------------------------------------------------------------
# Format-bazlı loader'lar
# ---------------------------------------------------------------------------

def load_pdf(path: Path, *, ocr_fallback: bool = True, ocr_lang: str = "fra") -> list[PageContent]:
    """
    pdfplumber ile sayfa-bazlı metin çıkarır.

    Bir sayfa boş döndüğünde (tarama / imaj-bazlı PDF), `ocr_fallback=True`
    ise sayfa pypdfium2 ile imaja render edilir ve pytesseract ile OCR'lanır.
    tesseract yoksa fallback sessizce atlanır; sayfa boş raporlanır.
    """
    try:
        import pdfplumber
    except ImportError as exc:
        raise LoaderError("pdfplumber yüklü değil; pip install pdfplumber") from exc

    pages: list[PageContent] = []
    with pdfplumber.open(str(path)) as pdf:
        total = len(pdf.pages)
        for i, page in enumerate(pdf.pages, start=1):
            raw = page.extract_text() or ""
            if not raw.strip() and ocr_fallback:
                try:
                    raw = _ocr_pdf_page(page, lang=ocr_lang)
                except LoaderError:
                    # tesseract eksik veya render hatası — sayfa boş kalır;
                    # smoke raporu durumu yine de yansıtır.
                    raw = ""
            pages.append(PageContent(text=_normalize(raw), page=i, total_pages=total))
    return pages


def _ocr_pdf_page(page, *, lang: str = "fra") -> str:
    """Pdf sayfasını imaja render edip OCR'lar. tesseract gerektirir."""
    try:
        import pytesseract
    except ImportError as exc:
        raise LoaderError("pytesseract yüklü değil") from exc

    try:
        # 300 DPI = tesseract'ın önerdiği asgari; 200 DPI küçük/renkli tablo
        # metnini (ör. niveaux_mots_cles.pdf slaytları) okuyamıyordu.
        img = page.to_image(resolution=300).original  # PIL Image (pypdfium2 backend)
        return _tesseract_image_to_string(img, lang=lang)
    except _tesseract_runtime_errors(pytesseract) as exc:
        raise LoaderError(f"OCR başarısız: {type(exc).__name__}: {exc}") from exc
    except Exception as exc:  # noqa: BLE001 — render/OCR backendleri farklı hatalar yükseltebilir
        raise LoaderError(f"OCR render başarısız: {type(exc).__name__}: {exc}") from exc


def load_docx(path: Path) -> list[PageContent]:
    """
    python-docx ile DOCX'ten metin + tablo hücrelerini birleştirir.

    DOCX'in kavramsal sayfası yoktur; tek PageContent döner (page=None).
    Tablo hücreleri dahil edilir çünkü Introduction_critères_corriges.docx
    gibi tablo-yoğun belgelerde içeriğin büyük kısmı tablodadır.
    """
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise LoaderError("python-docx yüklü değil; pip install python-docx") from exc

    doc = docx.Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    text = _normalize("\n".join(parts))
    return [PageContent(text=text, page=None, total_pages=1)]


def load_pptx(path: Path) -> list[PageContent]:
    """python-pptx ile slide-bazlı metin çıkarır; her slide bir PageContent."""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise LoaderError("python-pptx yüklü değil; pip install python-pptx") from exc

    prs = Presentation(str(path))
    total = len(prs.slides)
    pages: list[PageContent] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text)
        text = _normalize(" ".join(parts))
        pages.append(PageContent(text=text, page=i, total_pages=total))
    return pages


def load_image_ocr(path: Path, *, lang: str = "fra") -> list[PageContent]:
    """
    pytesseract ile bir imajı OCR eder; tek PageContent döner.

    tesseract sistem ikilisi veya French dil paketi (`tesseract-fra`) yoksa
    LoaderError yükseltir; smoke script bu durumu raporlar.
    """
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise LoaderError(
            "pytesseract veya Pillow yüklü değil; pip install pytesseract Pillow"
        ) from exc

    try:
        with Image.open(str(path)) as img:
            text = _tesseract_image_to_string(img, lang=lang, source_path=path)
    except _tesseract_runtime_errors(pytesseract) as exc:
        raise LoaderError(f"OCR başarısız: {type(exc).__name__}: {exc}") from exc

    return [PageContent(text=_normalize(text), page=None, total_pages=1)]


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

def load_document(path: Path) -> list[PageContent]:
    """Uzantıya göre uygun loader'ı çağırır."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext == ".pptx":
        return load_pptx(path)
    if ext in IMAGE_EXTENSIONS:
        return load_image_ocr(path)
    raise LoaderError(f"Desteklenmeyen uzantı: {ext}")


def iter_documents(root: Path) -> Iterator[LoadResult]:
    """
    `root` altındaki desteklenen tüm dosyaları özyinelemeli gez; her dosya
    için bir `LoadResult` yield et. Hatalar yakalanır ve `error` alanına
    yazılır; iterator akışı kesilmez.

    Gizli dosyalar (`.DS_Store` gibi) atlanır.
    """
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.name.startswith("."):
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            pages = load_document(path)
            yield LoadResult(path=path, pages=pages, error=None)
        except LoaderError as exc:
            logger.warning("Loader hatası %s: %s", path, exc)
            yield LoadResult(path=path, pages=[], error=str(exc))
        except Exception as exc:  # noqa: BLE001 — bozuk dosya / parser çökmesi
            logger.exception("Extraction beklenmedik biçimde başarısız: %s", path)
            yield LoadResult(path=path, pages=[], error=f"{type(exc).__name__}: {exc}")
